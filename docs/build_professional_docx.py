import os
import re
import sys
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
DOCS_DIR = os.path.join(BASE_DIR, "docs")
DIAG_DIR = os.path.join(DOCS_DIR, "diagrams")
SC_DIR = os.path.join(BASE_DIR, "screenshots")

CODE_SC_DIR = os.path.join(DOCS_DIR, "code_screenshots")

COLOR_BLACK = RGBColor(0, 0, 0)         # Pure Black for text on no background or light background
COLOR_WHITE = RGBColor(255, 255, 255)   # Pure White for text on dark background

def set_page_setup(doc):
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.0)
        
        # Header & Footer
        header = section.header
        p_hdr = header.paragraphs[0]
        p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_hdr.paragraph_format.space_after = Pt(0)
        r_hdr = p_hdr.add_run("SEKRETARIAT DPRD PROVINSI JAWA BARAT  •  PLATFORM HUDANG")
        r_hdr.font.name = "Arial"
        r_hdr.font.size = Pt(8)
        r_hdr.font.color.rgb = COLOR_BLACK
        
        footer = section.footer
        p_ftr = footer.paragraphs[0]
        p_ftr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_ftr.paragraph_format.space_after = Pt(0)
        r_ftr = p_ftr.add_run("Dokumen Resmi SPBE  |  Halaman ")
        r_ftr.font.name = "Arial"
        r_ftr.font.size = Pt(8.5)
        r_ftr.font.color.rgb = COLOR_BLACK
        add_page_number(r_ftr)

def add_page_number(run):
    fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
    run._r.append(fldSimple)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f"""
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        </w:tblBorders>
    """)
    tblPr.append(borders)

def set_callout_border(cell, color="1E3A8A", sz="28"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f"""
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    """)
    tcPr.append(borders)

def add_kop_dokumen(doc, doc_title, doc_subtitle, doc_no, doc_date, doc_ver, doc_class):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.3)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=160, bottom=160, left=200, right=200)
    set_callout_border(cell, color="000000", sz="36")
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    r1 = p.add_run("PEMERINTAH DAERAH PROVINSI JAWA BARAT\n")
    r1.font.name = "Arial"
    r1.font.size = Pt(11)
    r1.bold = True
    r1.font.color.rgb = COLOR_BLACK
    
    r2 = p.add_run("SEKRETARIAT DEWAN PERWAKILAN RAKYAT DAERAH\n")
    r2.font.name = "Arial"
    r2.font.size = Pt(13)
    r2.bold = True
    r2.font.color.rgb = COLOR_BLACK
    
    r3 = p.add_run("Jalan Diponegoro No. 27, Telepon (022) 7208282, Kota Bandung, Jawa Barat 40115\n")
    r3.font.name = "Arial"
    r3.font.size = Pt(8.5)
    r3.font.color.rgb = COLOR_BLACK

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(18)
    p_title.paragraph_format.space_after = Pt(3)
    r_t = p_title.add_run(doc_title)
    r_t.font.name = "Arial"
    r_t.font.size = Pt(15)
    r_t.bold = True
    r_t.font.color.rgb = COLOR_BLACK
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(16)
    r_s = p_sub.add_run(doc_subtitle)
    r_s.font.name = "Arial"
    r_s.font.size = Pt(10)
    r_s.bold = True
    r_s.font.color.rgb = COLOR_BLACK

    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    set_table_borders(meta_table, color="CBD5E1", sz="4")
    
    meta_data = [
        ("Nomor Dokumen", doc_no),
        ("Status / Versi Efektif", f"{doc_ver} (Disahkan per {doc_date})"),
        ("Klasifikasi Dokumen", doc_class),
        ("Unit Kerja Pengampu", "Bagian Pengawasan & Humas, Sekretariat DPRD Provinsi Jawa Barat")
    ]
    
    col_widths = [Inches(2.2), Inches(4.1)]
    for row_idx, (label, val) in enumerate(meta_data):
        row = meta_table.rows[row_idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = col_widths[0]
        c1.width = col_widths[1]
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, top=80, bottom=80, left=120, right=120)
        set_cell_margins(c1, top=80, bottom=80, left=120, right=120)
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)
        r_lbl = p0.add_run(label)
        r_lbl.font.name = "Arial"
        r_lbl.font.size = Pt(8.5)
        r_lbl.bold = True
        r_lbl.font.color.rgb = COLOR_BLACK
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        r_val = p1.add_run(val)
        r_val.font.name = "Arial"
        r_val.font.size = Pt(8.5)
        r_val.font.color.rgb = COLOR_BLACK

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(12)

def add_formatted_text(p, text, base_font="Arial", base_size=9.5, base_color=COLOR_BLACK, is_justify=True):
    if is_justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(4)
    
    pattern = re.compile(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)')
    tokens = pattern.split(text)
    
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            run = p.add_run(token[2:-2])
            run.bold = True
            run.font.name = base_font
            run.font.size = Pt(base_size)
            run.font.color.rgb = COLOR_BLACK
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            run = p.add_run(token[1:-1])
            run.italic = True
            run.font.name = base_font
            run.font.size = Pt(base_size)
            run.font.color.rgb = COLOR_BLACK
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            run = p.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(base_size - 0.5)
            run.font.color.rgb = COLOR_BLACK
        else:
            run = p.add_run(token)
            run.font.name = base_font
            run.font.size = Pt(base_size)
            run.font.color.rgb = COLOR_BLACK

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(13)
    run.bold = True
    run.font.color.rgb = COLOR_BLACK
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(11)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(11)
    run.bold = True
    run.font.color.rgb = COLOR_BLACK
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10)
    run.bold = True
    run.font.color.rgb = COLOR_BLACK
    return p

def add_heading_4(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(9.5)
    run.bold = True
    run.font.color.rgb = COLOR_BLACK
    return p

def add_image_box(doc, img_path, caption, width_in=6.1):
    if not os.path.exists(img_path):
        print(f"WARNING: Image not found: {img_path}")
        return
        
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(img_path, width=Inches(width_in))

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(12)
    run_cap = p_cap.add_run(caption)
    run_cap.font.name = "Arial"
    run_cap.font.size = Pt(8.5)
    run_cap.font.italic = True
    run_cap.bold = True
    run_cap.font.color.rgb = COLOR_BLACK

def add_alert_box(doc, alert_type, text):
    configs = {
        "note": ("CATATAN TEKNIS:", "2563EB", "EFF6FF"),
        "important": ("CATATAN PENTING:", "7C3AED", "FAF5FF"),
        "warning": ("PERINGATAN OPERASIONAL:", "D97706", "FFFBEB"),
        "tip": ("PETUNJUK EFISIENSI:", "059669", "ECFDF5"),
        "caution": ("PERHATIAN KHUSUS:", "DC2626", "FEF2F2")
    }
    title, border_color, bg_color = configs.get(alert_type.lower(), ("INFORMASI:", "2563EB", "EFF6FF"))
    
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.3)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
    set_callout_border(cell, color=border_color, sz="28")
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    
    r_title = p.add_run(title + " ")
    r_title.bold = True
    r_title.font.name = "Arial"
    r_title.font.size = Pt(9)
    r_title.font.color.rgb = COLOR_BLACK
    
    add_formatted_text(p, text, base_font="Arial", base_size=9, base_color=COLOR_BLACK, is_justify=True)

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(4)

def add_code_box(doc, code_text, lang=""):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.3)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
    set_callout_border(cell, color="64748B", sz="12")
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    
    run = p.add_run(code_text.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(8)
    run.font.color.rgb = COLOR_BLACK

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(4)

def add_math_box(doc, math_text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.3)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
    set_callout_border(cell, color="0284C7", sz="20")
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    
    run = p.add_run(math_text.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.bold = True
    run.font.color.rgb = COLOR_BLACK

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(4)

def render_table(doc, rows_data):
    if not rows_data:
        return
        
    num_cols = max(len(r) for r in rows_data)
    num_rows = len(rows_data)
    
    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color="CBD5E1", sz="4")
    
    total_width = 6.3
    col_widths = [total_width / num_cols] * num_cols
    if num_cols >= 3 and len(rows_data[0][0]) <= 8:
        col_widths[0] = 0.6
        rem = (total_width - 0.6) / (num_cols - 1)
        for c in range(1, num_cols):
            col_widths[c] = rem
            
    for r_idx, row in enumerate(rows_data):
        is_header = (r_idx == 0)
        table_row = table.rows[r_idx]
        
        if is_header:
            trPr = table_row._tr.get_or_add_trPr()
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
            
        for c_idx in range(num_cols):
            cell = table_row.cells[c_idx]
            cell.width = Inches(col_widths[c_idx])
            val = row[c_idx] if c_idx < len(row) else ""
            val = val.replace("<br>", "\n").replace("<br/>", "\n")
            
            if is_header:
                # Dark background -> white text!
                set_cell_background(cell, "1E3A8A")
                set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            else:
                # Light background -> black text!
                bg = "FFFFFF" if (r_idx % 2 == 1) else "F8FAFC"
                set_cell_background(cell, bg)
                set_cell_margins(cell, top=70, bottom=70, left=100, right=100)
                
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.12
            
            if is_header:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(val)
                run.bold = True
                run.font.name = "Arial"
                run.font.size = Pt(8.5)
                run.font.color.rgb = COLOR_WHITE
            else:
                if len(val) <= 10 and re.match(r'^[0-9A-Za-z\.\-\_ ]+$', val.strip()):
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                add_formatted_text(p, val, base_font="Arial", base_size=8.5, base_color=COLOR_BLACK, is_justify=False)

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(6)

# -------------------------------------------------------------------------------------------------
# BUILD SOP DOCUMENT
# -------------------------------------------------------------------------------------------------
def build_sop_docx():
    print("\n[BUILDING] HUDANG_SOP_Alur_Kerja.docx...")
    doc = docx.Document()
    set_page_setup(doc)
    
    add_kop_dokumen(
        doc,
        doc_title="STANDAR OPERASIONAL PROSEDUR (SOP)\nPENYELENGGARAAN E-AUDIENSI DAN PENGAWALAN ASPIRASI BERBASIS PLATFORM HUDANG",
        doc_subtitle="Pedoman Baku Pelayanan E-Audiensi Tatap Muka Virtual Anggota DPRD dan Pengawalan Tindak Lanjut OPD",
        doc_no="SOP/DPRD-JBR/HUDANG/2026/001",
        doc_date="15 September 2026",
        doc_ver="Edisi 2.0 (Standar PermenPAN-RB No. 35/2012 - Selaras Aplikasi)",
        doc_class="Dokumen Tata Laksana Pelayanan Publik Kedewanan"
    )

    md_path = os.path.join(DOCS_DIR, "SOP_ALUR_KERJA_PLATFORM_HUDANG.md")
    with open(md_path, "r", encoding="utf-8") as f:
        content = f.read()

    lines = content.splitlines()
    in_table = False
    table_rows = []
    in_alert = False
    alert_type = ""
    alert_lines = []
    in_code = False
    code_lines = []
    
    for idx, line in enumerate(lines):
        trimmed = line.strip()

        if trimmed.startswith("# ") and "STANDAR OPERASIONAL" in trimmed:
            continue
        if trimmed.startswith("### SEKRETARIAT"):
            continue

        if trimmed.startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
                continue
            else:
                in_code = False
                add_code_box(doc, "\n".join(code_lines))
                code_lines = []
                continue

        if in_code:
            code_lines.append(line)
            continue

        if trimmed.startswith("|") and trimmed.endswith("|"):
            cells = [c.strip() for c in trimmed.split("|")[1:-1]]
            if all(re.match(r"^:?-+:?$", c) for c in cells):
                continue
            table_rows.append(cells)
            in_table = True
            continue
        elif in_table:
            render_table(doc, table_rows)
            table_rows = []
            in_table = False

        alert_match = re.match(r"^>\s*\[!(NOTE|IMPORTANT|WARNING|TIP|CAUTION)\]", trimmed, re.IGNORECASE)
        if alert_match:
            in_alert = True
            alert_type = alert_match.group(1).lower()
            alert_lines = []
            continue
        elif in_alert:
            if trimmed.startswith(">"):
                alert_lines.append(trimmed[1:].strip())
                continue
            else:
                add_alert_box(doc, alert_type, " ".join(alert_lines))
                in_alert = False
                alert_lines = []

        # Heading 1
        if trimmed.startswith("## "):
            h_text = trimmed[3:].strip()
            add_heading_1(doc, h_text)
            continue

        # Heading 2
        if trimmed.startswith("### "):
            h_text = trimmed[4:].strip()
            add_heading_2(doc, h_text)
            
            # Bagan Alur Prosedur SOP (Fokus murni ke SOP, tanpa screenshot aplikasi)
            if re.match(r"^5\.1\b", h_text):
                add_image_box(doc, os.path.join(DIAG_DIR, "sop_diagram_1.png"), "Gambar 5.1: Diagram Alur Makro Pengelolaan Aspirasi dan E-Audiensi (End-to-End)")
                add_image_box(doc, os.path.join(DIAG_DIR, "sop_main_flow.png"), "Gambar 5.2: Peta Alur Standar Pelayanan Publik Sekretariat DPRD Jawa Barat")
            elif re.match(r"^5\.2\b", h_text):
                add_image_box(doc, os.path.join(DIAG_DIR, "sop_diagram_2.png"), "Gambar 5.3: Diagram Alur Pengajuan dan Persetujuan Jadwal E-Audiensi")
            elif re.match(r"^5\.3\b", h_text):
                add_image_box(doc, os.path.join(DIAG_DIR, "sop_diagram_3.png"), "Gambar 5.4: Diagram Alur Kerja Pelaksanaan E-Audiensi Virtual dan Perekaman Resmi")
            elif re.match(r"^5\.4\b", h_text):
                add_image_box(doc, os.path.join(DIAG_DIR, "sop_diagram_4.png"), "Gambar 5.5: Diagram Alur Pelacakan Empat Tahap Tindak Lanjut Perangkat Daerah (OPD)")
            continue

        # Heading 3
        if trimmed.startswith("#### "):
            h_text = trimmed[5:].strip()
            add_heading_3(doc, h_text)
            continue

        if trimmed == "---":
            continue

        # Narrative Paragraph & List
        if trimmed:
            p = doc.add_paragraph()
            if re.match(r"^(\d+(\.\d+)*\.|[a-zA-Z]\.)\s+", trimmed):
                p.paragraph_format.left_indent = Inches(0.28)
                p.paragraph_format.first_line_indent = Inches(-0.28)
            elif trimmed.startswith("- ") or trimmed.startswith("* "):
                trimmed = "–  " + trimmed[2:].strip()
                p.paragraph_format.left_indent = Inches(0.28)
                p.paragraph_format.first_line_indent = Inches(-0.28)
            
            add_formatted_text(p, trimmed, base_font="Arial", base_size=9.5, base_color=COLOR_BLACK, is_justify=True)

    if in_table:
        render_table(doc, table_rows)
    if in_alert:
        add_alert_box(doc, alert_type, " ".join(alert_lines))

    out_path = os.path.join(DOCS_DIR, "HUDANG_SOP_Alur_Kerja.docx")
    doc.save(out_path)
    print(f"[COMPLETED] Saved: {out_path} ({os.path.getsize(out_path):,} bytes)")

# -------------------------------------------------------------------------------------------------
# BUILD SPESIFIKASI TEKNIS DOCUMENT
# -------------------------------------------------------------------------------------------------
def build_spec_docx(use_code_screenshots=False, output_filename="HUDANG_Spesifikasi_Teknis.docx"):
    print(f"\n[BUILDING] {output_filename}...")
    doc = docx.Document()
    set_page_setup(doc)
    
    subtitle = "Acuan Baku Rekayasa Perangkat Lunak, Infrastruktur Komputasi, Skema Data, dan Standar Keamanan Sistem"
    if use_code_screenshots:
        subtitle += " (Edisi Visual Cuplikan Kode IDE)"
        
    add_kop_dokumen(
        doc,
        doc_title="DOKUMEN SPESIFIKASI TEKNIS DAN ARSITEKTUR SISTEM\nPLATFORM ASPIRASI DAN E-AUDIENSI BERBASIS APLIKASI (HUDANG)",
        doc_subtitle=subtitle,
        doc_no="SPEC/HUDANG/01/2026",
        doc_date="14 September 2026",
        doc_ver="Versi 2.0 (Produksi SPBE - Visual IDE Code)" if use_code_screenshots else "Versi 2.0 (Produksi SPBE)",
        doc_class="Dokumen Rekayasa Perangkat Lunak & Tata Kelola SPBE"
    )

    md_path = os.path.join(DOCS_DIR, "SPESIFIKASI_TEKNIS_PLATFORM_HUDANG.md")
    with open(md_path, "r", encoding="utf-8") as f:
        content = f.read()

    lines = content.splitlines()
    in_table = False
    table_rows = []
    in_alert = False
    alert_type = ""
    alert_lines = []
    in_code = False
    code_lang = ""
    code_lines = []
    in_math = False
    math_lines = []
    current_h = ""
    inserted_code_sc = set()
    
    for idx, line in enumerate(lines):
        trimmed = line.strip()

        if trimmed.startswith("# ") and "DOKUMEN SPESIFIKASI" in trimmed:
            continue
        if trimmed.startswith("### SEKRETARIAT"):
            continue

        # Code block handling
        if trimmed.startswith("```"):
            if not in_code:
                in_code = True
                code_lang = trimmed[3:].strip().lower()
                code_lines = []
                continue
            else:
                in_code = False
                # If mermaid, do not output raw code! It is handled by high-res diagram image!
                if code_lang == "mermaid":
                    pass
                elif not use_code_screenshots:
                    add_code_box(doc, "\n".join(code_lines), lang=code_lang)
                else:
                    # Code Screenshot Mode!
                    if re.match(r"^###\s*2\.3\b", current_h):
                        if "code_2_3_nginx" not in inserted_code_sc:
                            inserted_code_sc.add("code_2_3_nginx")
                            add_image_box(doc, os.path.join(CODE_SC_DIR, "code_2_3_nginx_part1.png"), "Gambar 2.4a: Cuplikan Konfigurasi Gateway Proksi Nginx (SSL, WAF & Security Headers)")
                            add_image_box(doc, os.path.join(CODE_SC_DIR, "code_2_3_nginx_part2.png"), "Gambar 2.4b: Cuplikan Konfigurasi Gateway Proksi Nginx (WebSocket, Media & Rate Limiting)")
                    elif re.match(r"^###\s*2\.5\b", current_h) and code_lang in ["yaml", "yml"]:
                        if "code_2_5_livekit" not in inserted_code_sc:
                            inserted_code_sc.add("code_2_5_livekit")
                            add_image_box(doc, os.path.join(CODE_SC_DIR, "code_2_5_livekit.png"), "Gambar 2.5: Cuplikan Konfigurasi Server Gateway Media LiveKit SFU (livekit.yaml)")
                    elif re.match(r"^####\s*2\.5\.1\b", current_h) and code_lang in ["conf", "ini"]:
                        if "code_2_5_turnserver" not in inserted_code_sc:
                            inserted_code_sc.add("code_2_5_turnserver")
                            add_image_box(doc, os.path.join(CODE_SC_DIR, "code_2_5_turnserver.png"), "Gambar 2.6: Cuplikan Konfigurasi Server NAT Coturn STUN/TURN (turnserver.conf)")
                    elif re.match(r"^###\s*2\.6\b", current_h):
                        if "code_2_6_sysctl" not in inserted_code_sc:
                            inserted_code_sc.add("code_2_6_sysctl")
                            add_image_box(doc, os.path.join(CODE_SC_DIR, "code_2_6_sysctl.png"), "Gambar 2.7: Cuplikan Konfigurasi Optimasi Kernel Socket Linux (99-hudang-performance.conf)")
                    elif re.match(r"^###\s*3\.1\b", current_h):
                        if "code_3_1_api_envelope" not in inserted_code_sc:
                            inserted_code_sc.add("code_3_1_api_envelope")
                            add_image_box(doc, os.path.join(CODE_SC_DIR, "code_3_1_api_envelope.png"), "Gambar 3.6: Cuplikan Standarisasi Format Kontrak Respons REST API (Success & Error)")
                    elif re.match(r"^####\s*3\.3\.1\b", current_h):
                        if "code_3_3_jwt" not in inserted_code_sc:
                            inserted_code_sc.add("code_3_3_jwt")
                            add_image_box(doc, os.path.join(CODE_SC_DIR, "code_3_3_jwt.png"), "Gambar 3.7: Cuplikan Skema Endpoint Autentikasi dan Struktur Klaim Token JWT")
                    elif re.match(r"^####\s*3\.3\.2\b", current_h):
                        if "code_3_3_livekit_token" not in inserted_code_sc:
                            inserted_code_sc.add("code_3_3_livekit_token")
                            add_image_box(doc, os.path.join(CODE_SC_DIR, "code_3_3_livekit_token.png"), "Gambar 3.8: Cuplikan Skema Endpoint Token Akses WebRTC LiveKit SFU")
                    elif re.match(r"^####\s*3\.3\.3\b", current_h):
                        if "code_3_3_sipd" not in inserted_code_sc:
                            inserted_code_sc.add("code_3_3_sipd")
                            add_image_box(doc, os.path.join(CODE_SC_DIR, "code_3_3_sipd.png"), "Gambar 3.9: Cuplikan Format Kamus Usulan Interoperabilitas SIPD Republik Indonesia")
                    elif re.match(r"^####\s*3\.3\.4\b", current_h):
                        if "code_3_3_geojson" not in inserted_code_sc:
                            inserted_code_sc.add("code_3_3_geojson")
                            add_image_box(doc, os.path.join(CODE_SC_DIR, "code_3_3_geojson.png"), "Gambar 3.10: Cuplikan Skema Spasial Tematik GeoJSON FeatureCollection (RFC 7946)")
                    elif re.match(r"^###\s*4\.1\b", current_h):
                        if "code_4_1_docker_compose" not in inserted_code_sc:
                            inserted_code_sc.add("code_4_1_docker_compose")
                            add_image_box(doc, os.path.join(CODE_SC_DIR, "code_4_1_docker_compose.png"), "Gambar 4.3: Cuplikan Kode Sumber Orkestrasi Kontainer Sistem (docker-compose.yml)")
                    elif re.match(r"^###\s*4\.2\b", current_h):
                        if "code_4_2_schema" not in inserted_code_sc:
                            inserted_code_sc.add("code_4_2_schema")
                            add_image_box(doc, os.path.join(CODE_SC_DIR, "code_4_2_schema_part1.png"), "Gambar 4.4a: Cuplikan Skema Relasional Prisma (Bagian 1: User, Aspirasi, Verification, Schedule)")
                            add_image_box(doc, os.path.join(CODE_SC_DIR, "code_4_2_schema_part2.png"), "Gambar 4.4b: Cuplikan Skema Relasional Prisma (Bagian 2: ActionFollowUp, Feedback, Egress, SessionLog)")
                    elif re.match(r"^###\s*4\.3\b", current_h):
                        if "code_4_3_auth_routes" not in inserted_code_sc:
                            inserted_code_sc.add("code_4_3_auth_routes")
                            add_image_box(doc, os.path.join(CODE_SC_DIR, "code_4_3_auth_routes.png"), "Gambar 4.5: Cuplikan Kode Sumber Otentikasi dan Pengamanan Akses API (auth.routes.ts)")
                    elif re.match(r"^###\s*4\.4\b", current_h):
                        if "code_4_4_analysis_service" not in inserted_code_sc:
                            inserted_code_sc.add("code_4_4_analysis_service")
                            add_image_box(doc, os.path.join(CODE_SC_DIR, "code_4_4_analysis_service.png"), "Gambar 4.6: Cuplikan Kode Sumber Pipeline Transkripsi & Risalah AI (analysisService.ts)")
                    elif re.match(r"^###\s*4\.5\b", current_h):
                        if "code_4_5_livekit_routes" not in inserted_code_sc:
                            inserted_code_sc.add("code_4_5_livekit_routes")
                            add_image_box(doc, os.path.join(CODE_SC_DIR, "code_4_5_livekit_routes.png"), "Gambar 4.7: Cuplikan Kode Sumber Pembuatan Token Konferensi Video (livekit.routes.ts)")
                    elif re.match(r"^###\s*4\.6\b", current_h):
                        if "code_4_6_portal_tsx" not in inserted_code_sc:
                            inserted_code_sc.add("code_4_6_portal_tsx")
                            add_image_box(doc, os.path.join(CODE_SC_DIR, "code_4_6_portal_tsx.png"), "Gambar 4.8: Cuplikan Kode Sumber Antarmuka Keterbukaan Publik (PublicTransparencyPortal.tsx)")
                    elif re.match(r"^###\s*4\.7\b", current_h):
                        if "code_4_7_queue_service" not in inserted_code_sc:
                            inserted_code_sc.add("code_4_7_queue_service")
                            add_image_box(doc, os.path.join(CODE_SC_DIR, "code_4_7_queue_service.png"), "Gambar 4.9: Cuplikan Kode Sumber Layanan Antrean Latar Belakang FFmpeg (queueService.ts)")
                    elif re.match(r"^###\s*6\.3\b", current_h):
                        if "code_6_3_k6_load_test" not in inserted_code_sc:
                            inserted_code_sc.add("code_6_3_k6_load_test")
                            add_image_box(doc, os.path.join(CODE_SC_DIR, "code_6_3_k6_load_test.png"), "Gambar 6.3: Cuplikan Kode Sumber Skrip Pengujian Beban dan Konkurensi (k6_load_test.js)")
                    else:
                        add_code_box(doc, "\n".join(code_lines), lang=code_lang)
                code_lines = []
                code_lang = ""
                continue

        if in_code:
            code_lines.append(line)
            continue

        # Display Math ($$)
        if trimmed.startswith("$$"):
            if not in_math:
                in_math = True
                math_lines = []
                rest = trimmed[2:].strip()
                if rest.endswith("$$") and len(rest) > 2:
                    add_math_box(doc, rest[:-2].strip())
                    in_math = False
                elif rest:
                    math_lines.append(rest)
                continue
            else:
                in_math = False
                add_math_box(doc, "\n".join(math_lines))
                math_lines = []
                continue

        if in_math:
            if trimmed.endswith("$$"):
                math_lines.append(trimmed[:-2].strip())
                add_math_box(doc, "\n".join(math_lines))
                in_math = False
                math_lines = []
            else:
                math_lines.append(trimmed)
            continue

        # Table handling
        if trimmed.startswith("|") and trimmed.endswith("|"):
            cells = [c.strip() for c in trimmed.split("|")[1:-1]]
            if all(re.match(r"^:?-+:?$", c) for c in cells):
                continue
            table_rows.append(cells)
            in_table = True
            continue
        elif in_table:
            render_table(doc, table_rows)
            table_rows = []
            in_table = False

        # Alert handling
        alert_match = re.match(r"^>\s*\[!(NOTE|IMPORTANT|WARNING|TIP|CAUTION)\]", trimmed, re.IGNORECASE)
        if alert_match:
            in_alert = True
            alert_type = alert_match.group(1).lower()
            alert_lines = []
            continue
        elif in_alert:
            if trimmed.startswith(">"):
                alert_lines.append(trimmed[1:].strip())
                continue
            else:
                add_alert_box(doc, alert_type, " ".join(alert_lines))
                in_alert = False
                alert_lines = []

        # Heading 1 (##)
        if trimmed.startswith("## "):
            h_text = trimmed[3:].strip()
            current_h = trimmed
            add_heading_1(doc, h_text)
            continue

        # Heading 2 (###)
        if trimmed.startswith("### "):
            h_text = trimmed[4:].strip()
            current_h = trimmed
            add_heading_2(doc, h_text)
            
            # Map diagrams & screenshots to exact headings in Spesifikasi Teknis
            if re.match(r"^1\.4\b", h_text):
                add_image_box(doc, os.path.join(SC_DIR, "landing_page.png"), "Gambar 1.1: Portal Publik Pelayanan Aspirasi dan Pelacakan Status Tiket")
                add_image_box(doc, os.path.join(SC_DIR, "masyarakat_page.png"), "Gambar 1.2: Dasbor Pengajuan Usulan Warga dan Manajemen Riwayat Aspirasi")
                add_image_box(doc, os.path.join(SC_DIR, "admin_page.png"), "Gambar 1.3: Antarmuka Desk Verifikasi Berkas ASN dan Disposisi Komisi DPRD")
                add_image_box(doc, os.path.join(SC_DIR, "dewan_page.png"), "Gambar 1.4: Dasbor Ketersediaan Jadwal dan Konfirmasi E-Audiensi Anggota Dewan")
                add_image_box(doc, os.path.join(SC_DIR, "profile_page.png"), "Gambar 1.5: Antarmuka Manajemen Profil Pengguna dan Pengaturan Keamanan Akun")
                add_image_box(doc, os.path.join(SC_DIR, "gis_page.png"), "Gambar 1.6: Visualisasi Spasial Peta Tematik Persebaran Aspirasi 27 Kabupaten/Kota")
                add_image_box(doc, os.path.join(SC_DIR, "gis_kunjungan_page.png"), "Gambar 1.7: Pemetaan Spasial Lokasi Kunjungan Kerja Lapangan Pengawasan Dewan")
                add_image_box(doc, os.path.join(SC_DIR, "admin_settings_page.png"), "Gambar 1.8: Dasbor Konfigurasi Parameter Sistem, Jam Kerja, dan Batas Waktu SLA")
            elif re.match(r"^2\.1\b", h_text):
                add_image_box(doc, os.path.join(DIAG_DIR, "spec_diagram_1.png"), "Gambar 2.1: Arsitektur Multi-Tier Sistem Platform HUDANG (Presentation, Core API, SFU, Persistence)")
            elif re.match(r"^5\.1\b", h_text):
                add_image_box(doc, os.path.join(DIAG_DIR, "spec_diagram_2.png"), "Gambar 5.1: Skema Diagram Entitas Basis Data Relasional (ERD) Platform HUDANG")
            elif re.match(r"^6\.4\b", h_text):
                add_image_box(doc, os.path.join(DIAG_DIR, "spec_pitr_backup.png"), "Gambar 6.1: Siklus Pemulihan dan Pencadangan Basis Data (Point-in-Time Recovery)")
            elif re.match(r"^6\.5\b", h_text):
                add_image_box(doc, os.path.join(DIAG_DIR, "spec_diagram_3.png"), "Gambar 6.2: Peta Jalan Iterasi Sprint Agile Platform HUDANG (2026 - 2028)")
            continue

        # Heading 3 (####)
        if trimmed.startswith("#### "):
            h_text = trimmed[5:].strip()
            current_h = trimmed
            add_heading_3(doc, h_text)
            
            if re.match(r"^2\.2\.1\b", h_text):
                add_image_box(doc, os.path.join(DIAG_DIR, "spec_topo_docker.png"), "Gambar 2.2: Diagram Topologi Jaringan Kontainer Docker dan Pemetaan Port Host Server")
            elif re.match(r"^2\.5\.1\b", h_text):
                add_image_box(doc, os.path.join(DIAG_DIR, "spec_ice_nat.png"), "Gambar 2.3: Diagram Alur Pengumpulan Kandidat ICE dan Penembusan Firewall NAT Coturn STUN/TURN")
            elif re.match(r"^3\.4\.1\b", h_text):
                add_image_box(doc, os.path.join(DIAG_DIR, "spec_auth_jwt.png"), "Gambar 3.1: Diagram Alur Autentikasi Identitas NIK dan Pengelolaan Token JWT")
                add_image_box(doc, os.path.join(SC_DIR, "login_page.png"), "Gambar 3.2: Antarmuka Autentikasi Pengguna Multi-Peran dengan Validasi Kredensial")
                add_image_box(doc, os.path.join(SC_DIR, "register_page.png"), "Gambar 3.3: Formulir Registrasi Akun Warga Berbasis Validasi NIK Kependudukan 16 Digit")
            elif re.match(r"^3\.4\.2\b", h_text):
                add_image_box(doc, os.path.join(SC_DIR, "room_prejoin_page.png"), "Gambar 3.4: Modul Uji Kelayakan Perangkat Kamera, Audio, dan Jaringan WebRTC")
            elif re.match(r"^3\.4\.3\b", h_text):
                add_image_box(doc, os.path.join(DIAG_DIR, "spec_transcription.png"), "Gambar 3.5: Pipeline Komputasi Transkripsi Suara Verbatim dan Ringkasan Risalah Gemini AI")
            elif re.match(r"^4\.7\.1\b", h_text):
                add_image_box(doc, os.path.join(DIAG_DIR, "spec_egress_hls.png"), "Gambar 4.1: Pipeline Komputasi Egress Recording, Transcoding, dan HLS Streaming")
                add_image_box(doc, os.path.join(SC_DIR, "admin_streaming_page.png"), "Gambar 4.2: Modul Konfigurasi Siaran Terbuka Live Streaming Sidang Publik")
            elif re.match(r"^5\.4\.1\b", h_text):
                add_image_box(doc, os.path.join(DIAG_DIR, "spec_fsm_ticket.png"), "Gambar 5.2: Diagram Mesin Status Siklus Hidup Tiket Audiensi (Finite State Machine)")
            continue

        if trimmed == "---":
            continue

        # Narrative Paragraph & List
        if trimmed:
            p = doc.add_paragraph()
            if re.match(r"^(\d+(\.\d+)*\.|[a-zA-Z]\.)\s+", trimmed):
                p.paragraph_format.left_indent = Inches(0.28)
                p.paragraph_format.first_line_indent = Inches(-0.28)
            elif trimmed.startswith("- ") or trimmed.startswith("* "):
                trimmed = "–  " + trimmed[2:].strip()
                p.paragraph_format.left_indent = Inches(0.28)
                p.paragraph_format.first_line_indent = Inches(-0.28)
            
            add_formatted_text(p, trimmed, base_font="Arial", base_size=9.5, base_color=COLOR_BLACK, is_justify=True)

    if in_table:
        render_table(doc, table_rows)
    if in_alert:
        add_alert_box(doc, alert_type, " ".join(alert_lines))

    out_path = os.path.join(DOCS_DIR, output_filename)
    doc.save(out_path)
    print(f"[COMPLETED] Saved: {out_path} ({os.path.getsize(out_path):,} bytes)")

if __name__ == "__main__":
    build_sop_docx()
    build_spec_docx(use_code_screenshots=False, output_filename="HUDANG_Spesifikasi_Teknis.docx")
    build_spec_docx(use_code_screenshots=True, output_filename="HUDANG_Spesifikasi_Teknis_Screenshot_Kode.docx")
    print("\n[SUCCESS] All DOCX documents generated successfully!")

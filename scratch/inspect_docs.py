import docx

doc = docx.Document('docs/HUDANG_SOP_Alur_Kerja.docx')
for i, p in enumerate(doc.paragraphs):
    if 'graphic' in p._p.xml:
        print(f'P{i}: text="{p.text[:60]}" has graphic')
for ti, t in enumerate(doc.tables):
    for r in t.rows:
        for c in r.cells:
            for p in c.paragraphs:
                if 'graphic' in p._p.xml:
                    print(f'T{ti} cell has graphic: {p.text[:40]}')

for rel_id, rel in doc.part.rels.items():
    if 'image' in rel.target_ref:
        print(f'Image rel: {rel.target_ref}')

import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def add_formatted_text(p, text, base_font="Arial", base_size=10, base_color=RGBColor(30, 41, 59), is_justify=True):
    if is_justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(4)
    
    # Tokenize by **bold**, *italic*, and `code`
    pattern = re.compile(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)')
    tokens = pattern.split(text)
    
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            run = p.add_run(token[2:-2])
            run.bold = True
            run.font.name = base_font
            run.font.size = Pt(base_size)
            run.font.color.rgb = RGBColor(15, 23, 42)
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            run = p.add_run(token[1:-1])
            run.italic = True
            run.font.name = base_font
            run.font.size = Pt(base_size)
            run.font.color.rgb = base_color
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            run = p.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(base_size - 0.5)
            run.font.color.rgb = RGBColor(30, 58, 138)
        else:
            run = p.add_run(token)
            run.font.name = base_font
            run.font.size = Pt(base_size)
            run.font.color.rgb = base_color

doc = docx.Document()
p = doc.add_paragraph()
add_formatted_text(p, "Ini adalah **teks tebal** dan *teks miring* serta `kode program` yang rapi.")
print("Paragraph runs:", len(p.runs))
for r in p.runs:
    print(f"  Run: '{r.text}' bold={r.bold} italic={r.italic} font={r.font.name}")

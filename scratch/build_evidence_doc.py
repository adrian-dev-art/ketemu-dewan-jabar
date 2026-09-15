import os, sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_code_block(doc, title, file_path, code_text, explanation=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    run_title = p.add_run(f"📌 Evidence Code: {title}")
    run_title.bold = True
    run_title.font.size = Pt(11)
    run_title.font.color.rgb = RGBColor(14, 116, 144)

    if file_path:
        p_path = doc.add_paragraph()
        p_path.paragraph_format.space_after = Pt(4)
        run_path = p_path.add_run(f"Path: {file_path}")
        run_path.italic = True
        run_path.font.size = Pt(9.5)
        run_path.font.color.rgb = RGBColor(100, 116, 139)

    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    p_code = cell.paragraphs[0]
    p_code.paragraph_format.space_before = Pt(0)
    p_code.paragraph_format.space_after = Pt(0)
    p_code.paragraph_format.line_spacing = 1.15
    run_code = p_code.add_run(code_text.strip())
    run_code.font.name = "Consolas"
    run_code.font.size = Pt(8.5)
    run_code.font.color.rgb = RGBColor(30, 41, 59)

    if explanation:
        p_exp = doc.add_paragraph()
        p_exp.paragraph_format.space_before = Pt(4)
        p_exp.paragraph_format.space_after = Pt(10)
        run_exp = p_exp.add_run(f"Penjelasan Evidence: {explanation}")
        run_exp.font.size = Pt(10)
        run_exp.font.color.rgb = RGBColor(51, 65, 85)

def add_screenshot(doc, title, img_path, caption=None):
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        run_t = p.add_run(f"🖼️ Screenshot Evidence: {title}")
        run_t.bold = True
        run_t.font.size = Pt(11)
        run_t.font.color.rgb = RGBColor(2, 132, 199)

        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(6.0))

        if caption:
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_after = Pt(10)
            run_cap = p_cap.add_run(f"Gambar: {caption}")
            run_cap.italic = True
            run_cap.font.size = Pt(9)
            run_cap.font.color.rgb = RGBColor(100, 116, 139)

def main():
    doc = docx.Document()
    
    # Page setup
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles setup
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Calibri'
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = RGBColor(30, 41, 59)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("DOKUMEN TAHAPAN PROYEK")
    r_title.bold = True
    r_title.font.size = Pt(22)
    r_title.font.color.rgb = RGBColor(15, 23, 42)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(20)
    r_sub = p_sub.add_run("Sistem Aspirasi & Video Conference Dewan (DPRD HUDANG)\nLaporan Evidence Implementasi Kode, Konfigurasi, & Pengujian")
    r_sub.font.size = Pt(13)
    r_sub.font.color.rgb = RGBColor(71, 85, 105)

    doc.save("Tahapan_Proyek.docx")
    print("Base document saved.")

if __name__ == "__main__":
    main()
